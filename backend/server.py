from fastapi import FastAPI, APIRouter, Depends, HTTPException, Header
from fastapi.responses import Response
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import io
import logging
import uuid
import re
import calendar
import httpx
import jwt
import bcrypt
import hmac
import hashlib
import secrets
import ipaddress
from html import escape
from html.parser import HTMLParser
from urllib.parse import urlparse
from pathlib import Path
from pydantic import BaseModel, Field, EmailStr
from typing import List, Optional, Literal
from datetime import datetime, timezone, timedelta

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

JWT_SECRET = os.environ.get('JWT_SECRET', 'dev-secret')
JWT_ALG = 'HS256'
EMERGENT_LLM_KEY = os.environ.get('EMERGENT_LLM_KEY', '')

# Emergent managed email (Resend). Base URL is a constant — never read from env so
# it survives deployment.
EMAIL_BASE_URL = "https://integrations.emergentagent.com"
EMAIL_KEY = os.environ.get('EMERGENT_EMAIL_KEY', '')
EMAIL_FROM_NAME = os.environ.get('EMAIL_FROM_NAME', 'Smart Home Finance')
EMAIL_REPLY_TO = os.environ.get('EMAIL_REPLY_TO')
OTP_PEPPER = os.environ.get('OTP_PEPPER', JWT_SECRET).encode('utf-8')

app = FastAPI(title="Smart Home Finance API")
api_router = APIRouter(prefix="/api")

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# ---------------- Helpers ----------------
def now_utc() -> datetime:
    return datetime.now(timezone.utc)

def hash_password(pw: str) -> str:
    return bcrypt.hashpw(pw.encode('utf-8'), bcrypt.gensalt()).decode('utf-8')

def verify_password(pw: str, hashed: str) -> bool:
    try:
        return bcrypt.checkpw(pw.encode('utf-8'), hashed.encode('utf-8'))
    except Exception:
        return False

def create_jwt(user_id: str) -> str:
    payload = {
        'user_id': user_id,
        'exp': now_utc() + timedelta(days=30),
        'iat': now_utc(),
    }
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALG)

def new_id(prefix: str = "id") -> str:
    return f"{prefix}_{uuid.uuid4().hex[:12]}"

def norm(s: Optional[str]) -> str:
    return re.sub(r'\s+', ' ', (s or '').strip().lower())

async def get_current_user(authorization: Optional[str] = Header(None)) -> dict:
    if not authorization or not authorization.startswith('Bearer '):
        raise HTTPException(status_code=401, detail="Missing token")
    token = authorization.split(' ', 1)[1]
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALG])
        user = await db.users.find_one({"user_id": payload['user_id']}, {"_id": 0, "password_hash": 0})
        if not user:
            raise HTTPException(status_code=401, detail="User not found")
        return user
    except jwt.PyJWTError:
        pass
    session = await db.user_sessions.find_one({"session_token": token}, {"_id": 0})
    if not session:
        raise HTTPException(status_code=401, detail="Invalid token")
    exp = session.get('expires_at')
    if exp and exp.tzinfo is None:
        exp = exp.replace(tzinfo=timezone.utc)
    if exp and exp < now_utc():
        raise HTTPException(status_code=401, detail="Session expired")
    user = await db.users.find_one({"user_id": session['user_id']}, {"_id": 0, "password_hash": 0})
    if not user:
        raise HTTPException(status_code=401, detail="User not found")
    return user

# ---------------- Models ----------------
class RegisterInput(BaseModel):
    email: EmailStr
    password: str
    name: str

class LoginInput(BaseModel):
    email: EmailStr
    password: str

class SessionInput(BaseModel):
    session_id: str

class ForgotInput(BaseModel):
    email: EmailStr

class ResetInput(BaseModel):
    email: EmailStr
    otp: str
    new_password: str

class UserOut(BaseModel):
    user_id: str
    email: str
    name: str
    picture: Optional[str] = None
    whatsapp: Optional[str] = None

class AuthResponse(BaseModel):
    token: str
    user: UserOut

class UserUpdate(BaseModel):
    name: Optional[str] = None
    whatsapp: Optional[str] = None
    picture: Optional[str] = None

class TransactionIn(BaseModel):
    type: Literal['income', 'expense']
    amount: float
    category: str
    title: str
    note: Optional[str] = ''
    date: Optional[str] = None
    child_id: Optional[str] = None
    link_type: Optional[str] = None  # 'bill' | 'shopping' | 'education' | 'savings'
    link_id: Optional[str] = None

class TransactionOut(TransactionIn):
    tx_id: str
    created_at: str

class ShoppingIn(BaseModel):
    name: str
    category: str
    budget: float
    realized: float = 0
    status: Literal['belum', 'sebagian', 'selesai'] = 'belum'
    frequency: str = 'Rutin'
    month: str

class ShoppingOut(ShoppingIn):
    item_id: str
    auto_realized: float = 0

class BillIn(BaseModel):
    name: str
    kind: Literal['rutin', 'cicilan', 'pinjaman', 'lainnya'] = 'rutin'
    category: str
    amount: float
    due_date: Optional[str] = None
    due_day: Optional[int] = None
    status: Literal['segera', 'belum', 'terlambat', 'lunas', 'ditangguhkan'] = 'belum'
    note: Optional[str] = ''

class BillOut(BillIn):
    bill_id: str
    auto_paid: bool = False
    paid_amount: float = 0

class SavingsIn(BaseModel):
    name: str
    target: float
    saved: float = 0
    color: Optional[str] = '#10D96A'
    icon: Optional[str] = 'piggy-bank'
    monthly: Optional[float] = 0
    note: Optional[str] = ''

class SavingsOut(SavingsIn):
    goal_id: str
    auto_saved: float = 0

class EducationChildIn(BaseModel):
    name: str
    school: str
    grade: str
    photo_url: Optional[str] = None

class EducationChildOut(EducationChildIn):
    child_id: str

class EducationItemIn(BaseModel):
    child_id: str
    name: str
    category: str
    budget: float
    realized: float = 0
    frequency: str = 'Bulanan'
    month: str
    status: Literal['belum', 'sebagian', 'lunas'] = 'belum'

class EducationItemOut(EducationItemIn):
    item_id: str
    auto_realized: float = 0

class ReportRequest(BaseModel):
    month: Optional[str] = None

# ---------------- Auth Routes ----------------
@api_router.post('/auth/register', response_model=AuthResponse)
async def register(body: RegisterInput):
    existing = await db.users.find_one({"email": body.email.lower()})
    if existing:
        raise HTTPException(status_code=400, detail="Email sudah terdaftar")
    user_id = new_id('user')
    user_doc = {
        "user_id": user_id,
        "email": body.email.lower(),
        "name": body.name,
        "picture": None,
        "whatsapp": None,
        "password_hash": hash_password(body.password),
        "provider": "password",
        "created_at": now_utc(),
    }
    await db.users.insert_one(user_doc)
    token = create_jwt(user_id)
    return AuthResponse(token=token, user=UserOut(user_id=user_id, email=body.email.lower(), name=body.name))

@api_router.post('/auth/login', response_model=AuthResponse)
async def login(body: LoginInput):
    user = await db.users.find_one({"email": body.email.lower()})
    if not user:
        raise HTTPException(status_code=404, detail="Akun belum terdaftar. Silakan daftar dulu.")
    if not user.get('password_hash'):
        raise HTTPException(status_code=400, detail="Akun ini terdaftar via Google. Gunakan tombol Masuk dengan Google.")
    if not verify_password(body.password, user['password_hash']):
        raise HTTPException(status_code=401, detail="Kata sandi salah. Gunakan 'Lupa sandi?' untuk mengatur ulang.")
    token = create_jwt(user['user_id'])
    return AuthResponse(token=token, user=UserOut(
        user_id=user['user_id'], email=user['email'], name=user['name'],
        picture=user.get('picture'), whatsapp=user.get('whatsapp')
    ))

@api_router.post('/auth/session', response_model=AuthResponse)
async def session_exchange(body: SessionInput):
    async with httpx.AsyncClient(timeout=15) as hc:
        r = await hc.get(
            "https://demobackend.emergentagent.com/auth/v1/env/oauth/session-data",
            headers={"X-Session-ID": body.session_id},
        )
    if r.status_code != 200:
        raise HTTPException(status_code=401, detail="Session tidak valid")
    data = r.json()
    email = (data.get('email') or '').lower()
    name = data.get('name') or email.split('@')[0]
    picture = data.get('picture')
    session_token = data.get('session_token')
    if not email or not session_token:
        raise HTTPException(status_code=401, detail="Data sesi tidak lengkap")

    existing = await db.users.find_one({"email": email})
    if existing:
        user_id = existing['user_id']
        await db.users.update_one({"user_id": user_id}, {"$set": {"name": name, "picture": picture}})
        whatsapp = existing.get('whatsapp')
    else:
        user_id = new_id('user')
        await db.users.insert_one({
            "user_id": user_id, "email": email, "name": name, "picture": picture,
            "whatsapp": None, "provider": "google", "created_at": now_utc(),
        })
        whatsapp = None
    await db.user_sessions.insert_one({
        "session_token": session_token, "user_id": user_id,
        "created_at": now_utc(), "expires_at": now_utc() + timedelta(days=7),
    })
    return AuthResponse(token=session_token, user=UserOut(
        user_id=user_id, email=email, name=name, picture=picture, whatsapp=whatsapp
    ))

@api_router.get('/auth/me', response_model=UserOut)
async def me(user=Depends(get_current_user)):
    return UserOut(
        user_id=user['user_id'], email=user['email'], name=user['name'],
        picture=user.get('picture'), whatsapp=user.get('whatsapp')
    )

@api_router.patch('/auth/me', response_model=UserOut)
async def update_me(body: UserUpdate, user=Depends(get_current_user)):
    upd = {k: v for k, v in body.dict().items() if v is not None}
    if upd:
        await db.users.update_one({"user_id": user['user_id']}, {"$set": upd})
    u = await db.users.find_one({"user_id": user['user_id']}, {"_id": 0, "password_hash": 0})
    return UserOut(
        user_id=u['user_id'], email=u['email'], name=u['name'],
        picture=u.get('picture'), whatsapp=u.get('whatsapp')
    )

@api_router.post('/auth/logout')
async def logout(authorization: Optional[str] = Header(None)):
    if authorization and authorization.startswith('Bearer '):
        token = authorization.split(' ', 1)[1]
        await db.user_sessions.delete_one({"session_token": token})
    return {"ok": True}

# ---------------- Password Reset (emailed OTP) ----------------
def otp_digest(otp: str) -> str:
    return hmac.new(OTP_PEPPER, otp.encode('utf-8'), hashlib.sha256).hexdigest()

# --- Email safety gate (defense-in-depth, per Resend playbook) ---
_SHORTENERS = ("bit.ly", "tinyurl.com", "t.co", "is.gd", "cutt.ly", "goo.gl", "rebrand.ly")
_CRED_ASK = ("reply with your password", "reply with the code", "send your password", "cvv",
             "send us your password", "enter your password below", "confirm your card number",
             "your full card number", "seed phrase", "recovery phrase", "verify your card",
             "social security number", "confirm your bank details")
_HOSTISH = re.compile(r"\b(?:https?://)?((?:[a-z0-9-]+\.)+[a-z]{2,})", re.I)

def _host_ok(host: str) -> bool:
    if not host or "xn--" in host:
        return False
    try:
        ipaddress.ip_address(host)
        return False
    except ValueError:
        pass
    return not any(host == s or host.endswith("." + s) for s in _SHORTENERS)

def _same_site(shown: str, real: str) -> bool:
    return shown == real or real.endswith("." + shown) or shown.endswith("." + real)

class _EmailScan(HTMLParser):
    def __init__(self):
        super().__init__()
        self.tags, self.urls, self.anchors = set(), [], []
        self._href, self._text = None, []
    def handle_starttag(self, tag, attrs):
        self.tags.add(tag.lower())
        self.urls += [v for k, v in attrs if k.lower() in ("href", "src") and v]
        if tag.lower() == "a":
            self._href = dict((k.lower(), v) for k, v in attrs).get("href")
            self._text = []
    def handle_data(self, data):
        if self._href is not None:
            self._text.append(data)
    def handle_endtag(self, tag):
        if tag.lower() == "a" and self._href is not None:
            self.anchors.append((self._href, "".join(self._text)))
            self._href, self._text = None, []

def _assert_safe_email(subject: str, html: str) -> None:
    scan = _EmailScan(); scan.feed(html)
    if scan.tags & {"form", "input", "textarea", "select"}:
        raise ValueError("No forms or input fields in email (G2)")
    body = f"{subject}\n{html}".lower()
    for p in _CRED_ASK:
        if p in body:
            raise ValueError(f"Email asks the recipient for credentials: {p!r} (G2)")
    for url in scan.urls:
        low = url.strip().lower()
        if low.startswith(("mailto:", "tel:", "cid:", "#")):
            continue
        if not low.startswith("https://"):
            raise ValueError(f"Email links/assets must be absolute https: {url!r} (G3)")
        host = urlparse(low).hostname or ""
        if not _host_ok(host) or urlparse(low).username is not None:
            raise ValueError(f"Shortened, numeric-host or credential-bearing URL: {url!r} (G3)")
    for href, text in scan.anchors:
        real = urlparse(href.strip().lower()).hostname or ""
        if not real:
            continue
        for m in _HOSTISH.finditer(text):
            if not _same_site(m.group(1).lower(), real):
                raise ValueError(f"Anchor text {m.group(1)!r} != real link host {real!r} (G3)")

async def send_email(*, to: str, subject: str, html: str) -> Optional[str]:
    _assert_safe_email(subject, html)
    payload = {"to": [to], "subject": subject, "html": html, "from_name": EMAIL_FROM_NAME}
    if EMAIL_REPLY_TO:
        payload["contact_email"] = EMAIL_REPLY_TO
    async with httpx.AsyncClient(timeout=30) as hc:
        resp = await hc.post(
            f"{EMAIL_BASE_URL}/api/v1/email/send",
            headers={"X-Email-Key": EMAIL_KEY},
            json=payload,
        )
    resp.raise_for_status()
    return resp.json().get("id")

async def send_reset_email(to: str, name: str, otp: str) -> None:
    subject = "Kode Atur Ulang Kata Sandi Smart Home Finance"
    html = (
        '<table role="presentation" width="100%"><tr><td style="padding:24px;'
        'font-family:Arial,sans-serif;color:#0F172A;max-width:480px">'
        '<h2 style="margin:0 0 8px">Atur Ulang Kata Sandi</h2>'
        f'<p style="margin:0 0 12px">Halo {escape(name)}, gunakan kode berikut untuk '
        'mengatur ulang kata sandi akun Smart Home Finance Anda.</p>'
        f'<p style="font-size:34px;font-weight:bold;letter-spacing:8px;color:#7C4DFF;'
        f'margin:16px 0">{escape(otp)}</p>'
        '<p style="margin:0 0 12px">Kode ini berlaku selama 10 menit dan hanya dapat '
        'digunakan satu kali.</p>'
        '<p style="font-size:12px;color:#888;margin-top:24px">Jika Anda tidak meminta ini, '
        'abaikan email ini. Dikirim oleh Smart Home Finance. Kami tidak pernah meminta kata '
        'sandi Anda melalui email.</p>'
        '</td></tr></table>'
    )
    await send_email(to=to, subject=subject, html=html)

@api_router.post('/auth/forgot-password')
async def forgot_password(body: ForgotInput):
    email = body.email.lower()
    user = await db.users.find_one({"email": email})
    if not user:
        raise HTTPException(status_code=404, detail="Email tidak terdaftar.")
    if not user.get('password_hash'):
        raise HTTPException(status_code=400, detail="Akun ini terdaftar via Google. Gunakan tombol Masuk dengan Google.")
    otp = f"{secrets.randbelow(1_000_000):06d}"
    now = now_utc()
    await db.password_resets.update_many(
        {"user_id": user['user_id'], "used": False}, {"$set": {"used": True}}
    )
    await db.password_resets.insert_one({
        "user_id": user['user_id'], "email": email,
        "otp_hash": otp_digest(otp), "created_at": now,
        "expires_at": now + timedelta(minutes=10), "used": False, "attempts": 0,
    })
    try:
        await send_reset_email(email, user.get('name') or email.split('@')[0], otp)
    except Exception as e:
        logger.error(f"Reset email failed: {e}")
        raise HTTPException(status_code=502, detail="Gagal mengirim email. Coba lagi nanti.")
    return {"ok": True, "message": "Kode verifikasi telah dikirim ke email Anda."}

@api_router.post('/auth/reset-password', response_model=AuthResponse)
async def reset_password(body: ResetInput):
    if len(body.new_password) < 6:
        raise HTTPException(status_code=400, detail="Kata sandi minimal 6 karakter")
    user = await db.users.find_one({"email": body.email.lower()})
    if not user:
        raise HTTPException(status_code=400, detail="Kode tidak valid atau sudah kedaluwarsa.")
    reset = await db.password_resets.find_one({
        "user_id": user['user_id'], "used": False,
        "expires_at": {"$gt": now_utc()}, "attempts": {"$lt": 5},
    }, sort=[("created_at", -1)])
    if not reset or not hmac.compare_digest(reset['otp_hash'], otp_digest(body.otp)):
        if reset:
            await db.password_resets.update_one({"_id": reset['_id']}, {"$inc": {"attempts": 1}})
        raise HTTPException(status_code=400, detail="Kode tidak valid atau sudah kedaluwarsa.")
    changed = await db.password_resets.update_one(
        {"_id": reset['_id'], "used": False}, {"$set": {"used": True}}
    )
    if changed.modified_count != 1:
        raise HTTPException(status_code=400, detail="Kode tidak valid atau sudah kedaluwarsa.")
    await db.users.update_one(
        {"user_id": user['user_id']}, {"$set": {"password_hash": hash_password(body.new_password)}}
    )
    token = create_jwt(user['user_id'])
    return AuthResponse(token=token, user=UserOut(
        user_id=user['user_id'], email=user['email'], name=user['name'],
        picture=user.get('picture'), whatsapp=user.get('whatsapp')
    ))

# ---------------- Transactions ----------------
def _clean_tx(doc: dict) -> dict:
    return {k: v for k, v in doc.items() if k not in ('_id', 'user_id', 'month')}

async def _linked_sum(user_id: str, link_type: str, link_id: str, month: Optional[str] = None, tx_type: Optional[str] = 'expense') -> float:
    """Sum of transactions explicitly allocated (linked) to a budget item."""
    q = {"user_id": user_id, "link_type": link_type, "link_id": link_id}
    if tx_type:
        q['type'] = tx_type
    txs = await db.transactions.find(q, {"_id": 0, "amount": 1, "month": 1}).to_list(5000)
    if month:
        txs = [t for t in txs if t.get('month') == month]
    return sum(float(t.get('amount') or 0) for t in txs)

async def _decorate_shopping(items: list, user_id: str) -> list:
    out = []
    for it in items:
        linked = await _linked_sum(user_id, 'shopping', it['item_id'], it.get('month'))
        realized = max(float(it.get('realized') or 0), linked)
        budget = float(it.get('budget') or 0)
        status = 'belum' if realized <= 0 else ('selesai' if realized >= budget else 'sebagian')
        out.append({**it, "realized": realized, "auto_realized": linked, "status": status})
    return out

async def _decorate_savings(items: list, user_id: str) -> list:
    out = []
    for g in items:
        linked = await _linked_sum(user_id, 'savings', g['goal_id'], None)
        saved = max(float(g.get('saved') or 0), linked)
        out.append({**g, "saved": saved, "auto_saved": linked})
    return out

@api_router.get('/transactions', response_model=List[TransactionOut])
async def list_transactions(user=Depends(get_current_user), month: Optional[str] = None, type: Optional[str] = None):
    q = {"user_id": user['user_id']}
    if type: q['type'] = type
    if month: q['month'] = month
    items = await db.transactions.find(q).sort([("date", -1), ("created_at", -1)]).to_list(500)
    return [_clean_tx(t) for t in items]

@api_router.post('/transactions', response_model=TransactionOut)
async def create_transaction(body: TransactionIn, user=Depends(get_current_user)):
    tx_id = new_id('tx')
    date_str = body.date or now_utc().isoformat()
    month = date_str[:7]
    doc = {
        "tx_id": tx_id, "user_id": user['user_id'],
        "type": body.type, "amount": body.amount,
        "category": body.category, "title": body.title,
        "note": body.note or '', "date": date_str, "month": month,
        "child_id": body.child_id,
        "link_type": body.link_type, "link_id": body.link_id,
        "created_at": now_utc().isoformat(),
    }
    await db.transactions.insert_one(doc.copy())
    return _clean_tx(doc)

@api_router.put('/transactions/{tx_id}', response_model=TransactionOut)
async def update_transaction(tx_id: str, body: TransactionIn, user=Depends(get_current_user)):
    date_str = body.date or now_utc().isoformat()
    month = date_str[:7]
    upd = {
        "type": body.type, "amount": body.amount, "category": body.category,
        "title": body.title, "note": body.note or '', "date": date_str,
        "month": month, "child_id": body.child_id,
        "link_type": body.link_type, "link_id": body.link_id,
    }
    res = await db.transactions.update_one({"tx_id": tx_id, "user_id": user['user_id']}, {"$set": upd})
    if res.matched_count == 0:
        raise HTTPException(404, "Not found")
    doc = await db.transactions.find_one({"tx_id": tx_id, "user_id": user['user_id']})
    return _clean_tx(doc)

@api_router.delete('/transactions/{tx_id}')
async def delete_transaction(tx_id: str, user=Depends(get_current_user)):
    res = await db.transactions.delete_one({"tx_id": tx_id, "user_id": user['user_id']})
    if res.deleted_count == 0:
        raise HTTPException(404, "Not found")
    return {"ok": True}

# ---------------- Shopping ----------------
@api_router.get('/shopping', response_model=List[ShoppingOut])
async def list_shopping(user=Depends(get_current_user), month: Optional[str] = None):
    q = {"user_id": user['user_id']}
    if month: q['month'] = month
    items = await db.shopping.find(q, {"_id": 0, "user_id": 0}).to_list(500)
    return await _decorate_shopping(items, user['user_id'])

@api_router.post('/shopping', response_model=ShoppingOut)
async def create_shopping(body: ShoppingIn, user=Depends(get_current_user)):
    item_id = new_id('shp')
    doc = {"item_id": item_id, "user_id": user['user_id'], **body.dict()}
    await db.shopping.insert_one(doc.copy())
    d = {k: v for k, v in doc.items() if k != 'user_id'}
    dec = await _decorate_shopping([d], user['user_id'])
    return dec[0]

@api_router.put('/shopping/{item_id}', response_model=ShoppingOut)
async def update_shopping(item_id: str, body: ShoppingIn, user=Depends(get_current_user)):
    await db.shopping.update_one({"item_id": item_id, "user_id": user['user_id']}, {"$set": body.dict()})
    doc = await db.shopping.find_one({"item_id": item_id, "user_id": user['user_id']}, {"_id": 0, "user_id": 0})
    if not doc: raise HTTPException(404, "Not found")
    dec = await _decorate_shopping([doc], user['user_id'])
    return dec[0]

@api_router.delete('/shopping/{item_id}')
async def delete_shopping(item_id: str, user=Depends(get_current_user)):
    res = await db.shopping.delete_one({"item_id": item_id, "user_id": user['user_id']})
    if res.deleted_count == 0: raise HTTPException(404, "Not found")
    return {"ok": True}

# ---------------- Bills with auto-status ----------------
async def _txs_expense_by_month(user_id: str) -> dict:
    """Returns { month: [tx,...] } for user's expense transactions."""
    txs = await db.transactions.find({"user_id": user_id, "type": "expense"}, {"_id": 0}).to_list(2000)
    out = {}
    for t in txs:
        out.setdefault(t.get('month', ''), []).append(t)
    return out

def _bill_month(b: dict) -> str:
    try:
        return b['due_date'][:7]
    except Exception:
        return ''

def _match_bill_tx(bill: dict, txs: list) -> bool:
    b_cat = norm(bill.get('category'))
    b_name = norm(bill.get('name'))
    b_amt = float(bill.get('amount') or 0)
    for t in txs:
        t_cat = norm(t.get('category'))
        t_title = norm(t.get('title'))
        t_amt = float(t.get('amount') or 0)
        cat_ok = (b_cat and (b_cat == t_cat or b_cat in t_cat or b_cat in t_title))
        name_ok = (b_name and (b_name in t_title or t_title in b_name))
        amt_ok = t_amt >= b_amt * 0.9  # tolerate 10% variance
        if (cat_ok or name_ok) and amt_ok:
            return True
    return False

def _bill_due_day(b: dict) -> int:
    dd = b.get('due_day')
    if dd:
        try:
            return max(1, min(31, int(dd)))
        except Exception:
            pass
    try:
        return int((b.get('due_date') or '')[8:10])
    except Exception:
        return 1

def _month_due_date(view_month: str, day: int):
    y = int(view_month[:4]); m = int(view_month[5:7])
    last = calendar.monthrange(y, m)[1]
    from datetime import date as _date
    return _date(y, m, min(day, last))

async def _decorate_bills(bills: list, user_id: str, view_month: str) -> list:
    """Recurring bills: one bill repeats every month. Status is computed per
    view month, 'lunas' only when an expense is allocated to it that month."""
    today = now_utc().date()
    current_month = now_utc().strftime('%Y-%m')
    out = []
    for b in bills:
        amt = float(b.get('amount') or 0)
        linked = await _linked_sum(user_id, 'bill', b.get('bill_id', ''), view_month)
        day = _bill_due_day(b)
        due = _month_due_date(view_month, day)
        auto_paid = False
        if amt > 0 and linked >= amt * 0.99:
            status = 'lunas'
            auto_paid = True
        elif view_month < current_month:
            status = 'terlambat'
        elif view_month > current_month:
            status = 'belum'
        else:
            if due < today:
                status = 'terlambat'
            elif (due - today).days <= 7:
                status = 'segera'
            else:
                status = 'belum'
        out.append({**b, "due_date": due.isoformat(), "due_day": day, "status": status, "auto_paid": auto_paid, "paid_amount": linked})
    return out

@api_router.get('/bills', response_model=List[BillOut])
async def list_bills(user=Depends(get_current_user), month: Optional[str] = None):
    view_month = month or now_utc().strftime('%Y-%m')
    items = await db.bills.find({"user_id": user['user_id']}, {"_id": 0, "user_id": 0}).to_list(500)
    return await _decorate_bills(items, user['user_id'], view_month)

@api_router.post('/bills', response_model=BillOut)
async def create_bill(body: BillIn, user=Depends(get_current_user)):
    bill_id = new_id('bill')
    data = body.dict()
    if not data.get('due_day') and data.get('due_date'):
        try:
            data['due_day'] = int(str(data['due_date'])[8:10])
        except Exception:
            data['due_day'] = 1
    doc = {"bill_id": bill_id, "user_id": user['user_id'], **data}
    await db.bills.insert_one(doc.copy())
    d = {k: v for k, v in doc.items() if k != 'user_id'}
    dec = await _decorate_bills([d], user['user_id'], now_utc().strftime('%Y-%m'))
    return dec[0]

@api_router.put('/bills/{bill_id}', response_model=BillOut)
async def update_bill(bill_id: str, body: BillIn, user=Depends(get_current_user)):
    data = body.dict()
    if not data.get('due_day') and data.get('due_date'):
        try:
            data['due_day'] = int(str(data['due_date'])[8:10])
        except Exception:
            data['due_day'] = 1
    await db.bills.update_one({"bill_id": bill_id, "user_id": user['user_id']}, {"$set": data})
    doc = await db.bills.find_one({"bill_id": bill_id, "user_id": user['user_id']}, {"_id": 0, "user_id": 0})
    if not doc: raise HTTPException(404, "Not found")
    dec = await _decorate_bills([doc], user['user_id'], now_utc().strftime('%Y-%m'))
    return dec[0]

@api_router.delete('/bills/{bill_id}')
async def delete_bill(bill_id: str, user=Depends(get_current_user)):
    res = await db.bills.delete_one({"bill_id": bill_id, "user_id": user['user_id']})
    if res.deleted_count == 0: raise HTTPException(404, "Not found")
    return {"ok": True}

# ---------------- Savings ----------------
@api_router.get('/savings', response_model=List[SavingsOut])
async def list_savings(user=Depends(get_current_user)):
    items = await db.savings.find({"user_id": user['user_id']}, {"_id": 0, "user_id": 0}).to_list(500)
    return await _decorate_savings(items, user['user_id'])

@api_router.post('/savings', response_model=SavingsOut)
async def create_savings(body: SavingsIn, user=Depends(get_current_user)):
    goal_id = new_id('goal')
    doc = {"goal_id": goal_id, "user_id": user['user_id'], **body.dict()}
    await db.savings.insert_one(doc.copy())
    d = {k: v for k, v in doc.items() if k != 'user_id'}
    dec = await _decorate_savings([d], user['user_id'])
    return dec[0]

@api_router.put('/savings/{goal_id}', response_model=SavingsOut)
async def update_savings(goal_id: str, body: SavingsIn, user=Depends(get_current_user)):
    await db.savings.update_one({"goal_id": goal_id, "user_id": user['user_id']}, {"$set": body.dict()})
    doc = await db.savings.find_one({"goal_id": goal_id, "user_id": user['user_id']}, {"_id": 0, "user_id": 0})
    if not doc: raise HTTPException(404, "Not found")
    dec = await _decorate_savings([doc], user['user_id'])
    return dec[0]

@api_router.delete('/savings/{goal_id}')
async def delete_savings(goal_id: str, user=Depends(get_current_user)):
    res = await db.savings.delete_one({"goal_id": goal_id, "user_id": user['user_id']})
    if res.deleted_count == 0: raise HTTPException(404, "Not found")
    return {"ok": True}

# ---------------- Education ----------------
@api_router.get('/education/children', response_model=List[EducationChildOut])
async def list_children(user=Depends(get_current_user)):
    items = await db.education_children.find({"user_id": user['user_id']}, {"_id": 0, "user_id": 0}).to_list(200)
    return items

@api_router.post('/education/children', response_model=EducationChildOut)
async def create_child(body: EducationChildIn, user=Depends(get_current_user)):
    child_id = new_id('child')
    doc = {"child_id": child_id, "user_id": user['user_id'], **body.dict()}
    await db.education_children.insert_one(doc.copy())
    return {k: v for k, v in doc.items() if k != 'user_id'}

@api_router.put('/education/children/{child_id}', response_model=EducationChildOut)
async def update_child(child_id: str, body: EducationChildIn, user=Depends(get_current_user)):
    res = await db.education_children.update_one(
        {"child_id": child_id, "user_id": user['user_id']}, {"$set": body.dict()}
    )
    if res.matched_count == 0: raise HTTPException(404, "Not found")
    doc = await db.education_children.find_one({"child_id": child_id, "user_id": user['user_id']}, {"_id": 0, "user_id": 0})
    return doc

@api_router.delete('/education/children/{child_id}')
async def delete_child(child_id: str, user=Depends(get_current_user)):
    await db.education_children.delete_one({"child_id": child_id, "user_id": user['user_id']})
    await db.education_items.delete_many({"child_id": child_id, "user_id": user['user_id']})
    return {"ok": True}

def _match_edu_tx(item: dict, txs: list) -> float:
    i_cat = norm(item.get('category'))
    i_name = norm(item.get('name'))
    i_child = item.get('child_id')
    total = 0.0
    for t in txs:
        t_cat = norm(t.get('category'))
        t_title = norm(t.get('title'))
        # Optional strict match: if tx has child_id and doesn't match, skip
        if t.get('child_id') and i_child and t['child_id'] != i_child:
            continue
        cat_ok = i_cat and (i_cat == t_cat or i_cat in t_cat)
        name_ok = i_name and (i_name in t_title or t_cat == i_name)
        if cat_ok or name_ok:
            total += float(t.get('amount') or 0)
    return total

async def _decorate_edu_items(items: list, user_id: str, view_month: str) -> list:
    """Recurring education items: repeat every month, realized only from
    transactions allocated to them within the view month."""
    out = []
    for it in items:
        auto = await _linked_sum(user_id, 'education', it['item_id'], view_month)
        realized = auto
        budget = float(it.get('budget') or 0)
        if realized <= 0:
            status = 'belum'
        elif realized >= budget:
            status = 'lunas'
        else:
            status = 'sebagian'
        out.append({**it, "realized": realized, "auto_realized": auto, "status": status})
    return out

@api_router.get('/education/items', response_model=List[EducationItemOut])
async def list_edu_items(user=Depends(get_current_user), child_id: Optional[str] = None, month: Optional[str] = None):
    view_month = month or now_utc().strftime('%Y-%m')
    q = {"user_id": user['user_id']}
    if child_id: q['child_id'] = child_id
    items = await db.education_items.find(q, {"_id": 0, "user_id": 0}).to_list(500)
    return await _decorate_edu_items(items, user['user_id'], view_month)

@api_router.post('/education/items', response_model=EducationItemOut)
async def create_edu_item(body: EducationItemIn, user=Depends(get_current_user)):
    item_id = new_id('eitem')
    doc = {"item_id": item_id, "user_id": user['user_id'], **body.dict()}
    await db.education_items.insert_one(doc.copy())
    d = {k: v for k, v in doc.items() if k != 'user_id'}
    dec = await _decorate_edu_items([d], user['user_id'], body.month or now_utc().strftime('%Y-%m'))
    return dec[0]

@api_router.put('/education/items/{item_id}', response_model=EducationItemOut)
async def update_edu_item(item_id: str, body: EducationItemIn, user=Depends(get_current_user)):
    res = await db.education_items.update_one(
        {"item_id": item_id, "user_id": user['user_id']}, {"$set": body.dict()}
    )
    if res.matched_count == 0: raise HTTPException(404, "Not found")
    doc = await db.education_items.find_one({"item_id": item_id, "user_id": user['user_id']}, {"_id": 0, "user_id": 0})
    dec = await _decorate_edu_items([doc], user['user_id'], body.month or now_utc().strftime('%Y-%m'))
    return dec[0]

@api_router.delete('/education/items/{item_id}')
async def delete_edu_item(item_id: str, user=Depends(get_current_user)):
    await db.education_items.delete_one({"item_id": item_id, "user_id": user['user_id']})
    return {"ok": True}

@api_router.get('/education/summary')
async def education_summary(user=Depends(get_current_user), month: Optional[str] = None):
    if not month:
        month = now_utc().strftime('%Y-%m')
    children = await db.education_children.find({"user_id": user['user_id']}, {"_id": 0, "user_id": 0}).to_list(200)
    items = await db.education_items.find({"user_id": user['user_id']}, {"_id": 0, "user_id": 0}).to_list(500)
    items = await _decorate_edu_items(items, user['user_id'], month)
    per_child = []
    total_budget = 0.0
    total_realized = 0.0
    for c in children:
        c_items = [i for i in items if i['child_id'] == c['child_id']]
        b = sum(float(i['budget']) for i in c_items)
        r = sum(float(i['realized']) for i in c_items)
        per_child.append({
            "child_id": c['child_id'], "name": c['name'], "school": c.get('school', ''),
            "grade": c.get('grade', ''), "photo_url": c.get('photo_url'),
            "budget": b, "realized": r, "item_count": len(c_items),
        })
        total_budget += b
        total_realized += r
    return {
        "month": month,
        "total_budget": total_budget,
        "total_realized": total_realized,
        "total_remaining": total_budget - total_realized,
        "child_count": len(children),
        "per_child": per_child,
    }

# ---------------- Dashboard Summary ----------------
@api_router.get('/dashboard/summary')
async def dashboard_summary(user=Depends(get_current_user), month: Optional[str] = None):
    if not month:
        month = now_utc().strftime('%Y-%m')
    txs = await db.transactions.find({"user_id": user['user_id'], "month": month}, {"_id": 0}).to_list(1000)
    income = sum(t['amount'] for t in txs if t['type'] == 'income')
    expense = sum(t['amount'] for t in txs if t['type'] == 'expense')
    balance = income - expense
    saving_rate = round((balance / income) * 100, 1) if income > 0 else 0.0

    # Cumulative (running) balance: all income - all expense up to and including this month.
    cum_txs = await db.transactions.find(
        {"user_id": user['user_id'], "month": {"$lte": month}},
        {"_id": 0, "amount": 1, "type": 1},
    ).to_list(100000)
    cum_income = sum(t['amount'] for t in cum_txs if t['type'] == 'income')
    cum_expense = sum(t['amount'] for t in cum_txs if t['type'] == 'expense')
    cumulative_balance = cum_income - cum_expense
    opening_balance = cumulative_balance - balance

    cat_map = {}
    for t in txs:
        if t['type'] == 'expense':
            cat_map[t['category']] = cat_map.get(t['category'], 0) + t['amount']
    total_exp = sum(cat_map.values()) or 1
    categories = [
        {"category": k, "amount": v, "percent": round(v / total_exp * 100, 1)}
        for k, v in sorted(cat_map.items(), key=lambda x: -x[1])
    ]

    cashflow = []
    base = datetime.strptime(month + '-01', '%Y-%m-%d')
    for i in range(5, -1, -1):
        y = base.year
        m = base.month - i
        while m <= 0:
            m += 12; y -= 1
        key = f"{y:04d}-{m:02d}"
        m_txs = await db.transactions.find({"user_id": user['user_id'], "month": key}, {"_id": 0}).to_list(1000)
        cashflow.append({
            "month": key,
            "income": sum(t['amount'] for t in m_txs if t['type'] == 'income'),
            "expense": sum(t['amount'] for t in m_txs if t['type'] == 'expense'),
        })

    available = opening_balance + income
    if income > 0:
        saving_rate = round((balance / income) * 100, 1)
    elif available > 0:
        saving_rate = round(((available - expense) / available) * 100, 1)
    else:
        saving_rate = 0.0

    if income == 0 and expense == 0 and cumulative_balance == 0:
        score = 0
    else:
        score = 50
        if available > 0:
            score = int(50 + ((available - expense) / available) * 50)
        if cumulative_balance < 0:
            score = min(score, 25)
        score = max(0, min(100, score))

    bills_raw = await db.bills.find({"user_id": user['user_id']}, {"_id": 0, "user_id": 0}).to_list(500)
    bills = await _decorate_bills(bills_raw, user['user_id'], month)
    upcoming = []
    today = now_utc().date()
    for b in bills:
        if b['status'] == 'lunas':
            continue
        try:
            d = datetime.fromisoformat(b['due_date']).date()
            days_left = (d - today).days
            if -30 <= days_left <= 30:
                upcoming.append({**b, "days_left": days_left})
        except Exception:
            pass
    upcoming.sort(key=lambda x: x['days_left'])

    savings_raw = await db.savings.find({"user_id": user['user_id']}, {"_id": 0, "user_id": 0}).to_list(200)
    savings = await _decorate_savings(savings_raw, user['user_id'])
    total_bills = sum(b['amount'] for b in bills if b['status'] != 'lunas')
    edu_items = await db.education_items.find({"user_id": user['user_id']}, {"_id": 0, "user_id": 0}).to_list(500)
    edu_total = sum(float(i.get('budget') or 0) for i in edu_items)
    shop_items = await db.shopping.find({"user_id": user['user_id'], "month": month}, {"_id": 0}).to_list(500)
    shop_total = sum(float(i.get('budget') or 0) for i in shop_items)
    bills_total = sum(b['amount'] for b in bills)
    obligations = {
        "tagihan": {"total": bills_total, "count": len(bills)},
        "pendidikan": {"total": edu_total, "count": len(edu_items)},
        "belanja": {"total": shop_total, "count": len(shop_items)},
        "total_commitment": bills_total + edu_total + shop_total,
    }

    return {
        "month": month, "income": income, "expense": expense, "balance": balance,
        "cumulative_balance": cumulative_balance, "opening_balance": opening_balance,
        "available": available,
        "saving_rate": saving_rate, "health_score": score,
        "expense_by_category": categories, "cashflow": cashflow,
        "upcoming_bills": upcoming[:5], "savings": savings,
        "total_bills": total_bills, "tx_count": len(txs),
        "obligations": obligations,
    }

# ---------------- AI Insight ----------------
@api_router.get('/insights/generate')
async def generate_insight(user=Depends(get_current_user), month: Optional[str] = None):
    if not month:
        month = now_utc().strftime('%Y-%m')
    txs = await db.transactions.find({"user_id": user['user_id'], "month": month}, {"_id": 0}).to_list(1000)
    income = sum(t['amount'] for t in txs if t['type'] == 'income')
    expense = sum(t['amount'] for t in txs if t['type'] == 'expense')
    prev = await db.transactions.find({"user_id": user['user_id'], "month": {"$lt": month}}, {"_id": 0, "amount": 1, "type": 1}).to_list(100000)
    opening = sum(t['amount'] for t in prev if t['type'] == 'income') - sum(t['amount'] for t in prev if t['type'] == 'expense')
    cumulative = opening + income - expense
    cat_map = {}
    for t in txs:
        if t['type'] == 'expense':
            cat_map[t['category']] = cat_map.get(t['category'], 0) + t['amount']

    if income == 0 and expense == 0 and opening == 0:
        return {"insight": "Belum ada data transaksi. Tambahkan pemasukan dan pengeluaran pertama Anda untuk mendapatkan analisis keuangan yang lebih akurat."}

    summary = (
        f"Bulan {month}: Saldo awal (akumulasi bulan-bulan sebelumnya) Rp{opening:,.0f}. "
        f"Pemasukan bulan ini Rp{income:,.0f}, Pengeluaran Rp{expense:,.0f}, "
        f"Selisih bulan ini Rp{income - expense:,.0f}, Saldo berjalan total Rp{cumulative:,.0f}. "
        f"Pengeluaran per kategori: "
        + ", ".join([f"{k} Rp{v:,.0f}" for k, v in cat_map.items()])
    )
    try:
        from emergentintegrations.llm.chat import LlmChat, UserMessage
        chat = LlmChat(
            api_key=EMERGENT_LLM_KEY,
            session_id=f"insight_{user['user_id']}_{month}",
            system_message=(
                "Anda adalah asisten keuangan keluarga Indonesia. "
                "Berikan analisis singkat 3-4 kalimat dalam Bahasa Indonesia yang ramah dan actionable. "
                "PENTING: nilai kesehatan keuangan dari SALDO BERJALAN total (bukan hanya selisih bulan ini). "
                "Jika pengeluaran bulan ini lebih besar dari pemasukan tetapi masih tertutup saldo bulan sebelumnya, "
                "jangan sebut 'minus' atau 'mengkhawatirkan' selama saldo berjalan masih positif. "
                "Sertakan satu rekomendasi konkret untuk menghemat. "
                "Gunakan format Rupiah (Rp). Jangan pakai markdown."
            ),
        ).with_model("anthropic", "claude-sonnet-5")
        reply = await chat.send_message(UserMessage(text=f"Analisa data keuangan berikut: {summary}"))
        return {"insight": reply}
    except Exception as e:
        logger.warning(f"AI insight failed: {e}")
        top_cat = max(cat_map.items(), key=lambda x: x[1]) if cat_map else None
        if cumulative >= 0:
            base = f"Saldo berjalan Anda sehat sebesar Rp{cumulative:,.0f}."
            if expense > income:
                base += f" Meski pengeluaran bulan ini (Rp{expense:,.0f}) lebih besar dari pemasukan (Rp{income:,.0f}), semuanya masih tertutup saldo bulan sebelumnya (Rp{opening:,.0f})."
            else:
                base += f" Bulan ini Anda surplus Rp{income - expense:,.0f}."
        else:
            base = f"Perhatian: saldo berjalan Anda minus Rp{abs(cumulative):,.0f} karena pengeluaran kumulatif melebihi pemasukan."
        if top_cat:
            base += f" Kategori '{top_cat[0]}' adalah pengeluaran terbesar bulan ini (Rp{top_cat[1]:,.0f}). Tinjau kategori ini untuk menghemat."
        return {"insight": base}

# ---------------- Reports ----------------
def _fmt_idr(n: float) -> str:
    try:
        return "Rp" + f"{int(round(n)):,}".replace(",", ".")
    except Exception:
        return f"Rp{n:.0f}"

async def _report_data(user_id: str, month: str) -> dict:
    txs = await db.transactions.find({"user_id": user_id, "month": month}, {"_id": 0}).sort([("date", 1)]).to_list(2000)
    income = sum(t['amount'] for t in txs if t['type'] == 'income')
    expense = sum(t['amount'] for t in txs if t['type'] == 'expense')
    cat_map = {}
    for t in txs:
        if t['type'] == 'expense':
            cat_map[t['category']] = cat_map.get(t['category'], 0) + t['amount']
    bills_raw = await db.bills.find({"user_id": user_id}, {"_id": 0, "user_id": 0}).to_list(500)
    bills = await _decorate_bills(bills_raw, user_id, month)
    prev = await db.transactions.find({"user_id": user_id, "month": {"$lt": month}}, {"_id": 0, "amount": 1, "type": 1}).to_list(100000)
    opening = sum(t['amount'] for t in prev if t['type'] == 'income') - sum(t['amount'] for t in prev if t['type'] == 'expense')
    return {
        "month": month, "income": income, "expense": expense,
        "balance": income - expense, "tx_count": len(txs),
        "opening_balance": opening, "ending_balance": opening + income - expense,
        "categories": sorted(cat_map.items(), key=lambda x: -x[1]),
        "transactions": [{k: t.get(k) for k in ("date", "type", "title", "category", "amount")} for t in txs],
        "bills": bills,
    }

async def _report_data_range(user_id: str, start: str, end: str) -> dict:
    txs_all = await db.transactions.find({"user_id": user_id}, {"_id": 0}).sort([("date", 1)]).to_list(5000)
    txs = [t for t in txs_all if start <= (t.get('date') or '')[:10] <= end]
    income = sum(t['amount'] for t in txs if t['type'] == 'income')
    expense = sum(t['amount'] for t in txs if t['type'] == 'expense')
    cat_map = {}
    for t in txs:
        if t['type'] == 'expense':
            cat_map[t['category']] = cat_map.get(t['category'], 0) + t['amount']
    bills_raw = await db.bills.find({"user_id": user_id}, {"_id": 0, "user_id": 0}).to_list(500)
    bills = await _decorate_bills(bills_raw, user_id, now_utc().strftime('%Y-%m'))
    before = [t for t in txs_all if (t.get('date') or '')[:10] < start]
    opening = sum(t['amount'] for t in before if t['type'] == 'income') - sum(t['amount'] for t in before if t['type'] == 'expense')
    return {
        "month": f"{start} s/d {end}", "income": income, "expense": expense,
        "balance": income - expense, "tx_count": len(txs),
        "opening_balance": opening, "ending_balance": opening + income - expense,
        "categories": sorted(cat_map.items(), key=lambda x: -x[1]),
        "transactions": [{k: t.get(k) for k in ("date", "type", "title", "category", "amount")} for t in txs],
        "bills": bills,
    }

async def _resolve_report(user_id: str, month: Optional[str], start: Optional[str], end: Optional[str]) -> dict:
    if start and end:
        return await _report_data_range(user_id, start, end)
    if not month:
        month = now_utc().strftime('%Y-%m')
    return await _report_data(user_id, month)

@api_router.get('/report/text')
async def report_text(user=Depends(get_current_user), month: Optional[str] = None, start: Optional[str] = None, end: Optional[str] = None):
    d = await _resolve_report(user['user_id'], month, start, end)
    label = d['month']
    lines = [
        f"📊 *LAPORAN KEUANGAN* — {label}",
        f"👤 {user.get('name', '')}",
        "",
        f"🏦 Saldo Awal: {_fmt_idr(d.get('opening_balance', 0))}",
        f"💰 Pemasukan: {_fmt_idr(d['income'])}",
        f"💸 Pengeluaran: {_fmt_idr(d['expense'])}",
        f"📈 Selisih Bulan Ini: {_fmt_idr(d['balance'])}",
        f"🧮 Saldo Akhir (Berjalan): {_fmt_idr(d.get('ending_balance', d['balance']))}",
        f"🧾 Jumlah transaksi: {d['tx_count']}",
    ]
    if d['categories']:
        lines.append("")
        lines.append("*Top Kategori Pengeluaran*")
        for k, v in d['categories'][:5]:
            lines.append(f"• {k}: {_fmt_idr(v)}")
    unpaid = [b for b in d['bills'] if b['status'] != 'lunas']
    if unpaid:
        lines.append("")
        lines.append("*Tagihan Belum Lunas*")
        for b in unpaid[:5]:
            lines.append(f"• {b['name']}: {_fmt_idr(b['amount'])}")
    lines.append("")
    lines.append("— Smart Home Finance")
    return {"text": "\n".join(lines)}

@api_router.get('/report/pdf')
async def report_pdf(user=Depends(get_current_user), month: Optional[str] = None, start: Optional[str] = None, end: Optional[str] = None):
    d = await _resolve_report(user['user_id'], month, start, end)
    label = d['month']
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors as rlcolors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title=f"Laporan {label}")
    styles = getSampleStyleSheet()
    h = ParagraphStyle('h', parent=styles['Heading1'], textColor=rlcolors.HexColor('#0F172A'))
    story = [
        Paragraph(f"Laporan Keuangan — {label}", h),
        Paragraph(f"Nama: {user.get('name','')}", styles['Normal']),
        Spacer(1, 12),
    ]
    summary_rows = [
        ["Saldo Awal", _fmt_idr(d.get('opening_balance', 0))],
        ["Pemasukan", _fmt_idr(d['income'])],
        ["Pengeluaran", _fmt_idr(d['expense'])],
        ["Selisih Bulan Ini", _fmt_idr(d['balance'])],
        ["Saldo Akhir (Berjalan)", _fmt_idr(d.get('ending_balance', d['balance']))],
        ["Jumlah Transaksi", str(d['tx_count'])],
    ]
    t = Table(summary_rows, colWidths=[220, 260])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, -1), rlcolors.HexColor('#F1F5F9')),
        ('BOX', (0, 0), (-1, -1), 0.5, rlcolors.grey),
        ('INNERGRID', (0, 0), (-1, -1), 0.25, rlcolors.grey),
        ('FONTSIZE', (0, 0), (-1, -1), 11),
        ('PADDING', (0, 0), (-1, -1), 8),
    ]))
    story.append(t)
    story.append(Spacer(1, 16))
    if d['transactions']:
        story.append(Paragraph("Detail Transaksi", styles['Heading2']))
        rows = [["Tanggal", "Jenis", "Kategori", "Judul", "Nominal"]]
        for tx in d['transactions']:
            rows.append([
                (tx.get('date') or '')[:10],
                "Masuk" if tx['type'] == 'income' else "Keluar",
                tx.get('category', ''), tx.get('title', ''),
                _fmt_idr(tx.get('amount', 0)),
            ])
        tb = Table(rows, colWidths=[70, 60, 90, 180, 100])
        tb.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), rlcolors.HexColor('#1F2937')),
            ('TEXTCOLOR', (0, 0), (-1, 0), rlcolors.white),
            ('GRID', (0, 0), (-1, -1), 0.25, rlcolors.grey),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('PADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(tb)
    doc.build(story)
    return Response(content=buf.getvalue(), media_type="application/pdf",
                    headers={"Content-Disposition": f"attachment; filename=Laporan-{label}.pdf"})

@api_router.get('/report/excel')
async def report_excel(user=Depends(get_current_user), month: Optional[str] = None, start: Optional[str] = None, end: Optional[str] = None):
    d = await _resolve_report(user['user_id'], month, start, end)
    label = d['month']
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "Ringkasan"
    ws.append([f"Laporan Keuangan — {label}"])
    ws['A1'].font = Font(bold=True, size=14)
    ws.append([f"Nama: {user.get('name','')}"])
    ws.append([])
    ws.append(["Saldo Awal", d.get('opening_balance', 0)])
    ws.append(["Pemasukan", d['income']])
    ws.append(["Pengeluaran", d['expense']])
    ws.append(["Selisih Bulan Ini", d['balance']])
    ws.append(["Saldo Akhir (Berjalan)", d.get('ending_balance', d['balance'])])
    ws.append(["Jumlah Transaksi", d['tx_count']])
    ws.column_dimensions['A'].width = 26
    ws.column_dimensions['B'].width = 22

    ws2 = wb.create_sheet("Transaksi")
    header = ["Tanggal", "Jenis", "Kategori", "Judul", "Nominal"]
    ws2.append(header)
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F2937")
    for tx in d['transactions']:
        ws2.append([
            (tx.get('date') or '')[:10],
            "Masuk" if tx['type'] == 'income' else "Keluar",
            tx.get('category', ''), tx.get('title', ''), tx.get('amount', 0),
        ])
    for col, width in zip("ABCDE", [14, 10, 18, 32, 16]):
        ws2.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return Response(content=buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename=Laporan-{label}.xlsx"})

# ---------------- App Wiring ----------------
app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.on_event("startup")
async def startup():
    await db.users.create_index("email", unique=True)
    await db.users.create_index("user_id", unique=True)
    await db.user_sessions.create_index("session_token", unique=True)
    await db.transactions.create_index([("user_id", 1), ("month", 1)])
    await db.bills.create_index("user_id")
    await db.shopping.create_index([("user_id", 1), ("month", 1)])
    await db.savings.create_index("user_id")
    await db.password_resets.create_index("expires_at", expireAfterSeconds=0)
    await db.password_resets.create_index([("user_id", 1), ("used", 1)])

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
